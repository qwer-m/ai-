"""只读核对已成功运行的真实产物，保存生产导出器原始 Excel 并逐格读回验收。"""

from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from openpyxl import load_workbook
from sqlalchemy import text

from core.db.database import SessionLocal
from core.db.model_defs import AgentRun
from modules.agent_platform.excel_export import EXPORT_COLUMNS, _project_case, build_test_cases_excel
from modules.agent_platform.results import persisted_test_generation_result
from modules.agent_platform.serialization import serialize_run
from modules.agent_platform.service import AgentPlatformService


def digest(value) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":"), default=str).encode("utf-8")).hexdigest()


def verify(run_id: int, xlsx_path: Path) -> dict:
    report = {"run_id": run_id, "observed_at": datetime.now(timezone.utc).isoformat(),
              "mysql_transaction": "READ ONLY", "model_requests": 0, "online_database_writes": 0,
              "http_transport_executed": False, "excel_created": False}
    with SessionLocal() as db:
        db.execute(text("SET TRANSACTION READ ONLY"))
        # 活跃运行只读取状态，不将中间候选当作最终产物。
        status = db.query(AgentRun.status).filter(AgentRun.id == run_id).scalar()
        report["run_status"] = status
        if status != "success":
            report["status"] = "waiting" if status in {"pending", "running", "waiting_approval"} else "unavailable"
            report["reason"] = "运行尚未成功，未创建或替换 Excel"
            db.rollback()
            return report
        run = db.get(AgentRun, run_id)
        source_hash = digest([run.input_payload, run.run_context, run.output_payload])
        artifact = deepcopy(persisted_test_generation_result(run))
        assert artifact and artifact["target_met"] and artifact["final_review"]["approved"] is True
        assert not artifact["final_review"]["differences"]
        cases = artifact["test_cases"]
        assert len(cases) == int(run.input_payload["case_budget"])
        assert len({case["case_id"] for case in cases}) == len(cases)
        serialized = serialize_run(run)
        assert serialized["test_generation_result"] == artifact
        export, export_status = AgentPlatformService(db).get_test_case_export_data(run_id=run_id, user_id=run.user_id)
        assert export_status == "found" and export["test_cases"] == cases
        # 调用产品现有导出器，只落盘其返回的文件，不重排数据或重做样式。
        excel_bytes = build_test_cases_excel(export["test_cases"])
        expected_rows = [[header for header, _ in EXPORT_COLUMNS], *[_project_case(case) for case in cases]]
        workbook = load_workbook(BytesIO(excel_bytes), read_only=False, data_only=False)
        try:
            assert workbook.sheetnames == ["测试用例"]
            sheet = workbook["测试用例"]
            assert sheet.max_row == len(cases) + 1 and sheet.max_column == len(EXPORT_COLUMNS)
            actual_rows = [["" if cell.value is None else cell.value for cell in row]
                           for row in sheet.iter_rows()]
            assert actual_rows == expected_rows
            assert all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row)
            assert sheet.freeze_panes == "A2"
            last_column = sheet.cell(row=1, column=len(EXPORT_COLUMNS)).column_letter
            assert sheet.auto_filter.ref == f"A1:{last_column}{len(cases) + 1}"
            assert all(cell.alignment.wrap_text for row in sheet.iter_rows(min_row=2, max_col=7) for cell in row)
            row_heights = {str(index): sheet.row_dimensions[index].height for index in range(2, len(cases) + 2)}
        finally:
            workbook.close()
        db.refresh(run)
        assert digest([run.input_payload, run.run_context, run.output_payload]) == source_hash
        assert not db.new and not db.dirty and not db.deleted
        db.rollback()
    # 文件只在终态、完整来源核对和内存读回全部通过后生成。
    if xlsx_path.exists():
        raise FileExistsError(f"禁止覆盖已有交付物: {xlsx_path}")
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    xlsx_path.write_bytes(excel_bytes)
    saved = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        saved_rows = [["" if value is None else value for value in row]
                      for row in saved["测试用例"].iter_rows(values_only=True)]
        assert saved_rows == expected_rows
    finally:
        saved.close()
    report.update({"status": "passed", "excel_created": True, "case_count": len(cases),
                   "checked_display_columns": [header for header, _ in EXPORT_COLUMNS],
                   "checked_data_cells": len(cases) * len(EXPORT_COLUMNS),
                   "all_display_cells_exact": True, "case_order_exact": True, "no_formula_cells": True,
                   "api_serializer_full_artifact_exact": True, "owned_export_service_full_cases_exact": True,
                   "source_unchanged": True, "source_snapshot_sha256": source_hash,
                   "artifact_sha256": digest(artifact), "case_content_sha256": digest(cases),
                   "display_rows_sha256": digest(expected_rows), "display_rows": expected_rows,
                   "row_heights": row_heights, "visual_render_executed": False,
                   "xlsx_path": str(xlsx_path.resolve()), "xlsx_bytes": len(excel_bytes),
                   "xlsx_sha256": hashlib.sha256(excel_bytes).hexdigest(),
                   "case_row_hashes": [{"case_id": case["case_id"], "source_sha256": digest(case),
                                         "display_sha256": digest(expected_rows[index + 1])}
                                        for index, case in enumerate(cases)]})
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-id", type=int, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    options = parser.parse_args()
    report = verify(options.run_id, options.xlsx)
    options.report.parent.mkdir(parents=True, exist_ok=True)
    options.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: value for key, value in report.items()
                      if key not in {"display_rows", "row_heights", "case_row_hashes"}}, ensure_ascii=False))


if __name__ == "__main__":
    main()
