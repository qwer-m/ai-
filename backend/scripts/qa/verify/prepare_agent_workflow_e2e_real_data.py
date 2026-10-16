"""用真实需求文档或历史运行准备完整工作流验收；默认只读，显式执行才创建运行。"""

from __future__ import annotations

import argparse
from collections import Counter
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys

BACKEND_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(BACKEND_ROOT))


def digest(value) -> str:
    return hashlib.sha256(json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str,
    ).encode("utf-8")).hexdigest()


def write_report(path: Path, report: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def source_evidence(db, run_id: int) -> dict:
    """摘要覆盖历史运行、全部节点和事件；不把事实正文重复写入预检报告。"""
    from core.db.model_defs import AgentRun, AgentNodeRun, AgentRunEvent

    records = {}
    for model, condition in (
        (AgentRun, AgentRun.id == run_id),
        (AgentNodeRun, AgentNodeRun.run_id == run_id),
        (AgentRunEvent, AgentRunEvent.run_id == run_id),
    ):
        ids = [value for (value,) in db.query(model.id).filter(condition).order_by(model.id).all()]
        values = []
        for identifier in ids:
            record = db.get(model, identifier)
            values.append(digest({column.key: getattr(record, column.key) for column in model.__table__.columns}))
        records[model.__tablename__] = {"count": len(ids), "ids": ids, "sha256": digest(values)}
    if records[AgentRun.__tablename__]["count"] != 1:
        raise ValueError(f"历史运行不存在，不能继续验证: run_id={run_id}")
    return records


def prepare(
    source_run_id: int | None = None, *, project_id: int | None = None,
    user_id: int | None = None, document_id: int | None = None,
    case_budget: int = 50, batch_case_limit: int = 5,
    context_compression_max_tokens: int = 1800,
) -> dict:
    from jsonschema import validate
    from sqlalchemy import text
    from core.db.database import SessionLocal
    from core.db.model_defs import AgentRun, AgentWorkflowDefinition
    from core.settings.config import settings
    from modules.agent_platform.contracts import AgentRunCreate, parse_execution_definition
    from modules.agent_platform.definition_repository import AgentDefinitionRepository
    from modules.agent_platform.lifecycle import ACTIVE_RUN_STATUSES, TERMINAL_RUN_STATUSES
    from modules.agent_platform.registry import BUILTIN_WORKFLOW_SPECS, runtime_registry_signature
    from modules.agent_platform.run_repository import AgentRunRepository
    from modules.agent_platform.service import _resolved_execution_limits
    from modules.agent_platform.sources import assert_same_source, persisted_source_snapshot
    from modules.system_components.repositories.project_repository import ProjectRepository

    with SessionLocal() as db:
        db.execute(text("SET TRANSACTION READ ONLY"))
        run = None
        if source_run_id is not None:
            if any(value is not None for value in (project_id, user_id, document_id)):
                raise ValueError("历史运行与新文档参数不能混用")
            run = db.get(AgentRun, source_run_id)
            if run is None or run.status != "success":
                raise ValueError("历史验收准备必须选择真实成功的运行")
            workflow = db.get(AgentWorkflowDefinition, run.workflow_definition_id)
            if workflow is None:
                raise ValueError("历史运行缺少工作流定义")
            project_id, user_id = int(run.project_id), int(run.user_id)
            request = AgentRunCreate(project_id=project_id, workflow_key=workflow.workflow_key,
                                     input_payload=deepcopy(run.input_payload))
        else:
            if any(value is None for value in (project_id, user_id, document_id)):
                raise ValueError("新文档验收必须指定项目、用户和需求文档")
            request = AgentRunCreate(
                project_id=project_id, workflow_key="test_generation",
                input_payload={
                    "requirement": "", "requirement_doc_id": document_id,
                    "case_budget": case_budget, "batch_case_limit": batch_case_limit,
                    "disable_result_cache": True, "enable_context_compression": True,
                    "context_compression_max_tokens": context_compression_max_tokens,
                },
            )
        if ProjectRepository(db).get_owned_project(project_id=project_id, user_id=user_id) is None:
            raise ValueError("项目不存在或不属于指定用户")
        spec = next(item for item in BUILTIN_WORKFLOW_SPECS if item["workflow_key"] == request.workflow_key)
        graph = parse_execution_definition(spec["definition"])
        validate(instance=request.input_payload, schema=graph.input_schema)
        repo = AgentRunRepository(db)
        current_source = repo.resolve_source_snapshot(project_id=project_id, input_payload=request.input_payload)
        if current_source is None:
            raise ValueError("完整验收需要真实需求来源快照")
        source = persisted_source_snapshot(run) if run is not None else current_source
        if source is None:
            raise ValueError("历史运行缺少来源快照")
        assert_same_source(source, current_source)
        workflow_ids = AgentDefinitionRepository(db).list_workflow_definition_ids(
            project_id=project_id, workflow_key=request.workflow_key,
        )
        rows = repo.list_run_sources(project_id=project_id, user_id=user_id,
                                    workflow_definition_ids=workflow_ids,
                                    statuses=ACTIVE_RUN_STATUSES | TERMINAL_RUN_STATUSES)
        scoped = [row for row in rows if row.source_key == source.key]
        terminal = [row for row in scoped if row.status in TERMINAL_RUN_STATUSES]
        active = [row for row in scoped if row.status in ACTIVE_RUN_STATUSES]
        retention_limit = len(terminal) + 1
        # 配置真实范围为 1..20；不能把 0 当作不存在的禁用开关。
        if retention_limit > 20:
            raise ValueError("所需历史保留数超过现有配置上限 20，拒绝创建可能删除证据的新运行")
        evidence = source_evidence(db, source_run_id) if run is not None else None
        report = {
            "status": "prepared" if not active else "blocked_by_active_run",
            "prepared_at": datetime.now(timezone.utc).isoformat(),
            "preparation_mysql_transaction": "READ ONLY", "new_run_started": False,
            "preparation_mode": "historical_run" if run is not None else "new_document",
            "source_run_id": source_run_id, "source_user_id": user_id,
            "source_snapshot": source.to_dict(), "source_unchanged": True,
            "source_evidence": evidence,
            "request": request.model_dump(exclude_none=True),
            "request_sha256": digest(request.model_dump(exclude_none=True)),
            "http_entry": "POST /api/agents/runs",
            "script_execution_entry": "AgentPlatformService.create_run -> fresh process -> run_agent_workflow",
            "http_authentication_and_celery_queue_covered": False,
            "current_graph": {"node_count": len(graph.nodes), "types": dict(Counter(n.node_type for n in graph.nodes)),
                              "sha256": digest(graph.model_dump()), "output_node_key": graph.output_node_key},
            "runtime_registry_signature": runtime_registry_signature(),
            "same_source_existing_runs": [{"run_id": row.run_id, "status": row.status} for row in scoped],
            "active_run_ids": [row.run_id for row in active],
            "worker_environment": {"AGENT_RUN_HISTORY_LIMIT": str(retention_limit)},
            "retention": {"current_process_limit": settings.AGENT_RUN_HISTORY_LIMIT,
                          "required_worker_limit": retention_limit, "existing_terminal_count": len(terminal),
                          "disabled_by_zero_supported": False},
            "cost_boundaries": {"quota_mode": "per_agent_instance",
                                "per_agent_instance_limits": _resolved_execution_limits(request),
                                "run_deadline_seconds": settings.AGENT_RUN_DEADLINE_SECONDS,
                                "aggregate_run_request_cap": None,
                                "historical_source_usage": deepcopy(run.run_context.get("usage") or {}) if run is not None else None,
                                "disable_result_cache": bool(request.input_payload.get("disable_result_cache"))},
            "execution_side_effects": [
                "同步内置全局 Agent、项目工具和工作流定义，停用已移除的内置定义",
                "创建新 AgentRun，保存节点、行动检查点、事件、模型用量和最终测试用例产物",
                "按真实配置调用模型；失败不转写为成功，不重置预算或自动创建第二个运行",
                "新进程按当前同源终态数量设置历史保留数，保留既有运行及其节点和事件",
            ],
        }
        assert not db.new and not db.dirty and not db.deleted
        db.rollback()
        return report


def worker(run_id: int, output: Path) -> None:
    """只由显式执行分支启动；环境在加载 settings 之前由父进程设置。"""
    from modules.agent_platform.runtime import run_agent_workflow

    result = run_agent_workflow(run_id=run_id)
    write_report(output, {"run_id": run_id, "runtime_result": result})


def execute(report: dict, output: Path) -> None:
    from core.db.database import SessionLocal
    from modules.agent_platform.contracts import AgentRunCreate
    from modules.agent_platform.service import AgentPlatformService
    from modules.agent_platform.registry import runtime_registry_signature
    from modules.agent_platform.sources import SourceSnapshot, assert_same_source

    if report["status"] != "prepared":
        raise ValueError("真实来源已有活动运行，本脚本不会接管或重复创建")
    worker_output = output.with_suffix(".worker.json")
    worker_log = output.with_suffix(".worker.log")
    if worker_output.exists() or worker_log.exists():
        raise ValueError("worker 输出路径已存在，拒绝覆盖已有证据")

    def start_current_worker(run_id: int):
        report.update(status="running", new_run_started=True, new_run_id=run_id,
                      worker_log=str(worker_log.resolve()), worker_result=str(worker_output.resolve()))
        write_report(output, report)
        worker_env = os.environ.copy()
        worker_env.update(report["worker_environment"])
        with worker_log.open("x", encoding="utf-8") as log:
            completed = subprocess.run(
                [sys.executable, str(Path(__file__).resolve()), "--worker-run-id", str(run_id),
                 "--output", str(worker_output.resolve())], cwd=BACKEND_ROOT, env=worker_env,
                stdout=log, stderr=subprocess.STDOUT, check=False,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        report["worker_process_exit_code"] = completed.returncode
        if completed.returncode:
            raise RuntimeError(f"独立真实执行器退出异常，详见本地 worker 日志: code={completed.returncode}")
        return json.loads(worker_output.read_text(encoding="utf-8"))

    with SessionLocal() as db:
        service = AgentPlatformService(db, start_current_worker)
        request = AgentRunCreate.model_validate(report["request"])
        current_source = service.runs.resolve_source_snapshot(
            project_id=request.project_id, input_payload=request.input_payload,
        )
        if current_source is None:
            raise ValueError("执行前缺少需求来源快照")
        assert_same_source(SourceSnapshot.from_dict(report["source_snapshot"]), current_source)
        if runtime_registry_signature() != report["runtime_registry_signature"]:
            raise ValueError("运行时代码在预检后变化，必须重新预检")
        run, reason = service.create_run(request=request,
                                         user_id=report["source_user_id"])
        report["create_result"] = reason
        if reason != "created" or run is None:
            raise ValueError(f"没有创建新的独立验收运行: {reason}")
        new_run_id = run.id

    verify_existing_run(report, new_run_id)


def verify_existing_run(report: dict, run_id: int) -> None:
    """只读复核真实终态；历史失败 attempt 作为证据保留，不冒充当前节点状态。"""
    from openpyxl import load_workbook
    from sqlalchemy import text
    from core.db.database import SessionLocal
    from core.db.model_defs import AgentRun, AgentWorkflowDefinition
    from modules.agent_platform.contracts import parse_execution_definition
    from modules.agent_platform.excel_export import build_test_cases_excel
    from modules.agent_platform.results import persisted_test_generation_result
    from modules.agent_platform.run_repository import AgentRunRepository
    from modules.agent_platform.serialization import serialize_run
    from modules.agent_platform.service import AgentPlatformService
    from modules.agent_platform.sources import SourceSnapshot, assert_same_source, persisted_source_snapshot

    if report.get("new_run_id") != run_id:
        raise ValueError("待验证运行与原始启动报告不一致")
    if digest(report["request"]) != report["request_sha256"]:
        raise ValueError("原始启动报告请求摘要不一致")
    with SessionLocal() as db:
        db.execute(text("SET TRANSACTION READ ONLY"))
        run = db.get(AgentRun, run_id)
        if run is None:
            raise ValueError("待验证的真实运行不存在")
        repo = AgentRunRepository(db)
        nodes = repo.list_node_runs(run_id=run_id)
        # 复用仓储既有的 attempt/id 选择规则，数据库仅排序轻量主键。
        latest = {key: repo.latest_node_run(run_id=run_id, node_key=key)
                  for key in {node.node_key for node in nodes}}
        workflow = db.get(AgentWorkflowDefinition, run.workflow_definition_id)
        graph = parse_execution_definition(workflow.definition) if workflow else None
        expected_keys = {node.node_key for node in graph.nodes} if graph else set()
        artifact = persisted_test_generation_result(run)
        unchanged = None
        if report.get("source_run_id") is not None:
            unchanged = source_evidence(db, report["source_run_id"]) == report["source_evidence"]
        expected_source = SourceSnapshot.from_dict(report["source_snapshot"])
        run_source = persisted_source_snapshot(run)
        current_source = repo.resolve_source_snapshot(project_id=run.project_id, input_payload=run.input_payload)
        if run_source is None or current_source is None:
            raise ValueError("正式运行或当前文档缺少需求来源快照")
        assert_same_source(expected_source, run_source)
        assert_same_source(expected_source, current_source)
        report.update(
            verification_mysql_transaction="READ ONLY", verification_model_requests_sent=0,
            source_evidence_unchanged=unchanged, source_snapshot_unchanged=True,
            final_run_status=run.status,
            actual_node_count=len(latest), node_statuses=dict(Counter(node.status for node in latest.values())),
            historical_attempt_count=len(nodes),
            historical_attempt_statuses=dict(Counter(node.status for node in nodes)),
            node_attempts=[{"node_run_id": node.id, "node_key": node.node_key,
                            "attempt": node.attempt, "status": node.status,
                            "is_latest": latest[node.node_key].id == node.id,
                            "error_type": node.error_message.split(":", 1)[0] if node.error_message else None,
                            "error_sha256": digest(node.error_message) if node.error_message else None,
                            "error_diagnostic": deepcopy((node.sdk_state or {}).get("error_diagnostic"))}
                           for node in nodes],
            runtime_error_type=run.error_message.split(":", 1)[0] if run.error_message else None,
            runtime_error_sha256=digest(run.error_message) if run.error_message else None,
            actual_usage=deepcopy(run.run_context.get("usage") or {}),
            artifact_sha256=digest(artifact) if artifact else None,
            artifact_case_count=len(artifact.get("test_cases") or []) if artifact else 0,
            final_review=deepcopy(artifact.get("final_review")) if artifact else None,
        )
        assert unchanged is not False, "历史源运行、节点或事件发生变化"
        assert run.project_id == report["request"]["project_id"] and run.user_id == report["source_user_id"]
        assert run.input_payload == report["request"]["input_payload"], "实际输入与原始启动请求不一致"
        assert graph and digest(graph.model_dump()) == report["current_graph"]["sha256"], "工作流定义与启动快照不一致"
        assert workflow.workflow_key == report["request"]["workflow_key"]
        assert run.status == "success", f"真实工作流未成功: {run.status}"
        assert set(latest) == expected_keys and len(latest) == report["current_graph"]["node_count"]
        assert all(node.status == "success" for node in latest.values()), "最新节点 attempt 未全部成功"
        assert artifact and artifact["target_met"] and artifact["final_review"]["approved"] is True
        assert not artifact["final_review"]["differences"]
        assert len(artifact["test_cases"]) == report["request"]["input_payload"]["case_budget"]
        serialized = serialize_run(run, node_runs=nodes)
        assert serialized["test_generation_result"] == artifact
        export, export_status = AgentPlatformService(db).get_test_case_export_data(
            run_id=run_id, user_id=report["source_user_id"],
        )
        assert export_status == "found" and export["test_cases"] == artifact["test_cases"]
        excel_bytes = build_test_cases_excel(export["test_cases"])
        workbook = load_workbook(io.BytesIO(excel_bytes), read_only=True, data_only=True)
        try:
            exported_ids = [row[0] for row in workbook.active.iter_rows(min_row=2, values_only=True)]
            assert exported_ids == [case["case_id"] for case in artifact["test_cases"]]
        finally:
            workbook.close()
        report["delivery_checks"] = {
            "api_serializer_matches_persisted_artifact": True,
            "owned_export_service_matches_persisted_cases": True,
            "excel_case_ids_and_order_match": True,
            "excel_bytes_sha256": hashlib.sha256(excel_bytes).hexdigest(),
            "http_transport_executed": False,
        }
        report["status"] = "passed"
        assert not db.new and not db.dirty and not db.deleted
        db.rollback()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-run-id", type=int)
    parser.add_argument("--project-id", type=int, help="新文档验收所属项目")
    parser.add_argument("--user-id", type=int, help="项目所属用户")
    parser.add_argument("--document-id", type=int, help="已真实解析成功的需求文档")
    parser.add_argument("--case-budget", type=int, default=50, help="新文档验收目标用例数，默认 50")
    parser.add_argument("--batch-case-limit", type=int, default=5, help="新文档验收每批用例上限，默认 5")
    parser.add_argument("--context-compression-max-tokens", type=int, default=1800,
                        help="新文档验收上下文压缩上限，默认 1800")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--execute", action="store_true", help="显式创建并执行正式新 Run；默认只读准备")
    parser.add_argument("--verify-run-id", type=int, help="只读复核已有运行，不创建 Run 或调用模型")
    parser.add_argument("--source-report", type=Path, help="原始启动报告；沿用启动前的历史证据摘要")
    parser.add_argument("--worker-run-id", type=int, help=argparse.SUPPRESS)
    options = parser.parse_args()
    new_document_values = (options.project_id, options.user_id, options.document_id)
    has_new_document = any(value is not None for value in new_document_values)
    if has_new_document:
        if any(value is None for value in new_document_values):
            parser.error("新文档验收必须同时提供 --project-id、--user-id 和 --document-id")
        if any(value is not None for value in (options.source_run_id, options.verify_run_id, options.worker_run_id)):
            parser.error("新文档验收参数不能与历史运行、只读复核或 worker 参数混用")
    output = options.output.resolve()
    if output.exists():
        raise ValueError("输出路径已存在，拒绝覆盖已有证据")
    if options.verify_run_id is not None:
        if options.execute or options.source_run_id is not None or options.worker_run_id is not None:
            parser.error("只读复核不能与创建或 worker 参数混用")
        if options.source_report is None:
            parser.error("只读复核必须提供 --source-report")
        source_report = options.source_report.resolve()
        source_bytes = source_report.read_bytes()
        report = json.loads(source_bytes)
        report.update(
            status="verifying", verification_mode="existing_run_read_only",
            source_execution_report=str(source_report),
            source_execution_report_sha256=hashlib.sha256(source_bytes).hexdigest(),
            source_execution_report_status=report.get("status"),
            source_execution_report_error_type=report.pop("error_type", None),
            verification_started_at=datetime.now(timezone.utc).isoformat(),
        )
        try:
            verify_existing_run(report, options.verify_run_id)
        except Exception as error:
            report.update(status="failed", error_type=type(error).__name__)
            raise
        finally:
            report["verification_finished_at"] = datetime.now(timezone.utc).isoformat()
            write_report(output, report)
        print(json.dumps({key: report.get(key) for key in (
            "status", "new_run_id", "final_run_status", "actual_node_count",
            "historical_attempt_count", "source_evidence_unchanged", "verification_model_requests_sent",
        )}, ensure_ascii=False))
        return
    if options.source_report is not None:
        parser.error("--source-report 仅用于 --verify-run-id")
    if options.worker_run_id is not None:
        if options.execute or options.source_run_id is not None:
            raise ValueError("worker 参数不能与创建参数混用")
        worker(options.worker_run_id, output)
        return
    if options.source_run_id is None and not has_new_document:
        parser.error("必须提供 --source-run-id，或同时提供 --project-id、--user-id 和 --document-id")
    report = prepare(
        options.source_run_id, project_id=options.project_id, user_id=options.user_id,
        document_id=options.document_id, case_budget=options.case_budget,
        batch_case_limit=options.batch_case_limit,
        context_compression_max_tokens=options.context_compression_max_tokens,
    )
    write_report(output, report)
    if options.execute:
        try:
            execute(report, output)
        except Exception as error:
            # 具体上游错误由正式 Runtime 诊断保存；这里不复制可能含认证信息的异常正文。
            report.update(status="failed", error_type=type(error).__name__)
            raise
        finally:
            report["finished_at"] = datetime.now(timezone.utc).isoformat()
            write_report(output, report)
    print(json.dumps({key: report.get(key) for key in (
        "status", "source_run_id", "new_run_started", "new_run_id", "current_graph", "worker_environment",
    )}, ensure_ascii=False))


if __name__ == "__main__":
    main()
